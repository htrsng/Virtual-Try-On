"""Biên soạn báo cáo từ bằng chứng đã có; không tự chạy hoặc sửa kết quả test."""
from collections import Counter
from datetime import datetime
from pathlib import Path
import json
import shutil

ROOT = Path(__file__).resolve().parents[2]


def read(relative):
    return json.loads((ROOT / relative).read_text(encoding='utf-8'))


def compile_report():
    campaign = read('results/spec_completed_48/report.json')
    comparison = read('results/comparison_3seeds/report.json')
    replay = read('results/reproduced_48/report.json')
    assert campaign['counts'] == {'PASS': 27, 'FAIL': 21}
    assert replay['counts'] == {'REPRODUCED': 21}
    pages = []
    pages.append('''## 1. Giới thiệu và mục tiêu

### 1.1. Bối cảnh

API bán hàng tiếp nhận dữ liệu do client cung cấp: thông tin đăng nhập, giỏ hàng, tổng tiền, trạng thái đơn và hồ sơ người dùng. Khi các ràng buộc không được kiểm tra tại server, một request hợp lệ về cú pháp vẫn có thể tạo ra dữ liệu sai nghiệp vụ. Đề tài lựa chọn fuzzing để khảo sát nhiều biến thể đầu vào và dùng oracle để nhận biết vi phạm.

### 1.2. Mục tiêu

- Xây dựng đặc tả và testcase truy vết được cho năm module bán hàng.
- Sinh input bằng seed, giá trị biên, sai kiểu và các phép đột biến JSON.
- Thực hiện fuzzing hộp đen và lựa chọn input theo phản hồi độ bao phủ.
- Lưu bằng chứng, chạy lại phát hiện từ trạng thái sạch và so sánh chiến lược với cùng ngân sách.

### 1.3. Phạm vi

Đối tượng là node_harness với ứng dụng Express, datastore trong bộ nhớ và các dịch vụ mô phỏng. Request thực sự được thực thi; kết quả không phải dữ liệu giả được điền sẵn. Kết luận chỉ áp dụng cho hành vi được harness mô phỏng, không suy rộng thành mức an toàn của hệ thống triển khai thực tế.

Nội dung chính tập trung vào fuzzing, tương ứng kỹ thuật Chương 4 của học phần. Những sheet CTL, C/Python và SMT được giữ theo mẫu đặc tả, nhưng 50 testcase minh họa không được tính vào kết quả kiểm thử API trong báo cáo này. Không khẳng định đã kiểm chứng toàn bộ các chương chỉ vì workbook có đủ sheet.
''')
    loc = len((ROOT / 'node_harness/index.js').read_text(encoding='utf-8').splitlines())
    pages.append(f'''## 2. Mô tả hệ thống mục tiêu

### 2.1. Thành phần và luồng dữ liệu

Python tạo input và điều phối; Node xử lý request trên ứng dụng mô phỏng; oracle so sánh phản hồi và dữ liệu cuối với điều kiện mong đợi. File index.js có {loc} dòng vật lý tại thời điểm biên soạn, gồm cả các route ngoài phạm vi. Coverage không đại diện cho toàn bộ server thật.

```text
Excel + seed → sinh kịch bản/đột biến JSON
                         ↓
                 HTTP hoặc Node harness
                         ↓
                 Express → datastore mô phỏng
                         ↓
              response + dữ liệu cuối + coverage
                         ↓
                  oracle → kết quả và log
```

Hình 1. Luồng dữ liệu thực nghiệm. Blackbox không dùng coverage để chọn input.

### 2.2. Bề mặt đầu vào

| Module | Endpoint chính | Dữ liệu/rủi ro được kiểm tra |
|---|---|---|
| Đăng nhập | POST /api/auth/login | Email, password, role trong token, bộ đếm |
| Tạo đơn | POST /api/orders | Giỏ hàng, số lượng, giá, tổng tiền, mã giảm giá |
| Trạng thái đơn | PUT /api/orders/:id | ID, quyền, enum, vòng đời, version |
| Người dùng | PUT /api/users/:id; POST /api/auth/create-admin | Quyền, role, mật khẩu, version, secret |
| Sản phẩm | POST /api/products; PUT /api/products/:id | Giá, tồn kho, ID trùng, quyền, version |

Bảng 1. Phạm vi API. Không phải mọi endpoint đều là target của cả hai fuzzer: blackbox có 7 target, whitebox có 6 target; đăng nhập trong whitebox phục vụ setup.
''')
    pages.append('''## 3. Phương pháp thực hiện

### 3.1. Từ đặc tả đến testcase

Năm workbook Module_*_spec.xlsx chứa yêu cầu chức năng, phi chức năng, ràng buộc, mô hình và testcase. ID yêu cầu là khóa truy vết để liên hệ kết quả với hành vi mong đợi. Đặc tả không được sửa cho khớp hành vi sai chỉ để testcase đạt.

Có 98 testcase tài liệu: 48 API và 50 minh họa. Bộ chạy đã triển khai đủ 48 testcase API. Các kịch bản nằm trong testcases.py, có ID, steps, actor và check. spec_campaign.py đọc workbook và điều phối các kế hoạch đó; nó chưa tự suy luận mọi yêu cầu bằng ngôn ngữ tự nhiên thành mã kiểm thử.

| Ví dụ yêu cầu | Input/kịch bản | Điều kiện đánh giá |
|---|---|---|
| Số lượng nguyên dương | Âm, sai kiểu, giỏ sai cấu trúc | Từ chối theo mã HTTP quy định |
| Tổng tiền theo catalog | Giá giả tự khớp tổng giả | Từ chối hoặc tính lại theo catalog |
| Cập nhật có version | Hai request cùng version | Một thành công, một 409; dữ liệu khớp request thắng |
| Bộ đếm nguyên tử | Ban đầu 7; 50 lần sai đồng thời | 50 phản hồi 401 và bộ đếm 57 |
| Trạng thái kết thúc | Hủy rồi yêu cầu giao | Từ chối và giữ dữ liệu sau hủy |

Bảng 2. Quan hệ yêu cầu, input và oracle. Đây là ví dụ thiết kế; trạng thái thực tế nằm ở phần kết quả.

16 testcase bổ sung được mở rộng thành 34 kịch bản, trong đó từng giá trị version sai được chạy độc lập. Cùng 62 kịch bản của bộ trước tạo tổng 96 kịch bản với cấu hình fuzz-cases=8. Các phép đo hiệu năng và rate limiting chưa được chứng minh bằng campaign này.
''')
    pages.append('''### 3.2. Sinh dữ liệu và fuzzing hộp đen

Seed là JSON hợp lệ làm điểm xuất phát. mutate_json() sao chép seed, chọn một phép đột biến rồi thay đổi dữ liệu. Bộ từ điển gồm null, chuỗi rỗng/dài, số âm/lớn, boolean, object, array và chuỗi có hình dạng tấn công. Có chuỗi SQL/XSS trong từ điển không có nghĩa đã xác nhận SQL injection hay XSS.

Các phép biến đổi là thay giá trị trường, xóa trường, đổi kiểu, sửa trường lồng và nhân bản phần tử mảng. Seed ngẫu nhiên được lưu để tái tạo chuỗi lựa chọn. Seed cố định không bảo đảm thời gian thực thi hay mọi dữ liệu sinh bởi môi trường luôn giống nhau.

Blackbox đăng nhập và tạo đối tượng chuẩn bị, sau đó luân phiên 7 target. Một số lượt bật/tắt token để khảo sát đường thiếu xác thực. Nó gửi HTTP, ghi mã phản hồi và áp dụng các quy tắc dùng chung trong fuzzer_common.py.

Dữ liệu tồn tại trong cả chiến dịch. Do đó input sửa sản phẩm có thể ảnh hưởng những request sau. Đây là tính chất cần ghi trong báo cáo: một input không phải lúc nào cũng tái hiện độc lập nếu thiếu trạng thái trước đó. Công cụ lưu riêng danh mục bất thường và ghi rõ phép đối chiếu nào không thể thực hiện khi danh mục đã hỏng.

```powershell
node fuzzing-security-testing/node_harness/http_server.js
```

Terminal thứ hai có thể chạy blackbox với --base-url http://127.0.0.1:5050. Khi không chỉ định URL, blackbox tự mở harness riêng. Cả hai cách vẫn dùng môi trường mô phỏng.
''')
    pages.append('''### 3.3. Fuzzing theo phản hồi độ bao phủ

whitebox_fuzzer.py sử dụng Node Inspector và v8-to-istanbul để thu các dòng đã thực thi của index.js. Inspector cung cấp cơ chế tương tác với bộ công cụ đo của runtime; tài liệu API được tham khảo ở [6]. File này hiện triển khai lựa chọn theo coverage, chưa triển khai thực thi tượng trưng hoặc giải ràng buộc concolic.

Mỗi target có hàng đợi riêng. Bộ chạy lấy input cha từ cuối hàng đợi, đột biến tối đa 6 lần (ENERGY=6), đo coverage từng input và giữ input con nếu có dòng mới. Hết hàng đợi thì quay lại seed; tổng input bị giới hạn đúng ngân sách.

```python
lines = set(result.get("covered_lines", []))
new = lines - covered
covered.update(lines)
if new:
    queue.append(child)
```

Đoạn mã 1. Nguyên tắc giữ input tạo coverage mới, trích từ fuzz_target(). Coverage chỉ hỗ trợ chọn đường khám phá; oracle vẫn cần để nhận biết sai nghiệp vụ.

Mỗi input chạy trong một tiến trình Node mới, vì vậy có dữ liệu seed sạch và phép đo riêng. Cách này thuận lợi cho tái hiện nhưng tốn chi phí khởi động. Coverage gồm cả setup/khởi tạo; số dòng lớn không tự chứng minh đã khám phá nhiều nhánh nghiệp vụ quan trọng.

Blackbox HTTP giữ trạng thái, còn whitebox đặt lại trạng thái. Để khảo sát lợi ích riêng của coverage feedback, bộ so sánh chạy cả baseline và feedback qua cùng harness cô lập, thay vì so trực tiếp hai cấu hình độc lập khác nhau.
''')
    pages.append('''### 3.4. Oracle, bằng chứng và giao thức thực thi

Oracle là điều kiện dùng để đánh giá hành vi quan sát được. Trong campaign, oracles.py kiểm tra status, token, bộ đếm, version, tổng tiền, trạng thái cuối và lưu mật khẩu. Với hai fuzzer, fuzzer_common.py cung cấp các kiểm tra response dùng chung.

Phép kiểm tra stale version không chỉ nhìn HTTP 409: dữ liệu sau request bị từ chối phải bằng dữ liệu trước đó. Phép kiểm tra đơn hợp lệ đối chiếu user, sản phẩm, số lượng và tổng đã lưu. Bộ đếm phải tăng tương đối từ giá trị ban đầu, không được reset về số request mới.

| Trạng thái/loại lỗi | Cách hiểu |
|---|---|
| PASS | Đã chạy và thỏa điều kiện trên các mẫu đó |
| FAIL | Có bằng chứng nhưng vi phạm oracle |
| ERROR | Setup, tiến trình hoặc đánh giá lỗi; chưa kết luận nghiệp vụ |
| PARTIAL | Chưa chạy đủ kế hoạch, chưa gặp FAIL/ERROR |
| NOT_RUN / NOT_IMPLEMENTED | Chưa chạy / chưa có kế hoạch thực thi |
| HTTP_5XX | Response lỗi ứng dụng, không tự chứng minh tiến trình crash |

Bảng 3. Quy ước thống kê. Finding là một lượt phát hiện; nhiều lượt có thể cùng nguyên nhân.

runtime.py kiểm tra dấu phân cách và định dạng JSON kết quả, định danh môi trường, mã thoát và timeout. JSON được ghi vào file tạm trước khi thay kết quả cũ. report.json lưu mã băm code; spec_snapshot.json lưu mã băm workbook. Công cụ tái hiện từ chối chạy nếu nguồn khác phiên bản campaign.
''')
    table = ['| Module | PASS | FAIL |', '|---|---|---|']
    for module in ['AUTH', 'ORD', 'OS', 'USR', 'PRD']:
        counts = Counter(e['status'] for e in campaign['traceability'] if e['testcase']['ID'].split('-')[1] == module)
        table.append(f"| {module} | {counts['PASS']} | {counts['FAIL']} |")
    pages.append('''## 4. Kết quả và phân tích

### 4.1. Campaign theo đặc tả

Nguồn chính là results/spec_completed_48/report.json. Cấu hình seed=7, fuzz-cases=8; môi trường node_harness_stub. Kết quả: 48 testcase, 96 kịch bản, 27 testcase PASS và 21 FAIL; theo kịch bản có 55 PASS, 41 FAIL và 0 ERROR. Tỷ lệ testcase đạt là 56,25%, chỉ mô tả mẫu đã chạy, không phải tỷ lệ an toàn của toàn hệ thống.

''' + '\n'.join(table) + '''

Bảng 4. Kết quả theo module. Testcase có nhiều kịch bản được đánh giá FAIL nếu có kịch bản vi phạm; không dùng số kịch bản làm mẫu số cho tỷ lệ testcase.

Bộ đếm từ 7 đạt 57 sau 50 request sai. Các testcase version bổ sung đạt; điều này không đồng nghĩa cơ chế tương thích bỏ version đã chống được stale write. Trong 16 testcase mới có 14 PASS, 2 FAIL: giá/tổng giả tự khớp và chuyển sang giao sau hủy.

Hai lần fuzzing độc lập trước đó ghi nhận blackbox 3500 input, 240 lượt lỗi, 1669 finding; whitebox 270 input, 31 lượt lỗi, 142 finding. Đây là bằng chứng lịch sử riêng, khác ngân sách và cách quản lý dữ liệu. Không dùng số finding lớn hơn để khẳng định phương pháp tốt hơn; không cộng chúng vào 48 testcase.

Kết quả Excel riêng nằm tại results/spec_completed_48/KetQuaKiemThu.xlsx. Đặc tả chỉ ghi điều kiện mong đợi; kết quả mỗi lần chạy được giữ trong results/ để truy vết.
''')
    pages.append('''### 4.2. Tái hiện và ví dụ vi phạm

Một kịch bản FAIL đại diện được chọn cho từng testcase rồi chạy lại hai lần từ Node mới. Cả 21 đại diện đều tái hiện FAIL, tổng 42 lượt, không ERROR. Kết quả được nhóm thành 13 nhóm dấu hiệu để điều tra, chưa xác định 13 nguyên nhân gốc độc lập.

**Ví dụ 1 — tin giá client:** TC-ORD-18 dùng productId=1 có giá catalog 200000; request gửi price=1, quantity=2, totalAmount=2, phí/giảm giá bằng 0. Ứng dụng trả thành công và lưu tổng 2. Oracle mong từ chối hoặc tính lại 400000. Vi phạm này ảnh hưởng tính toàn vẹn giá đơn; ưu tiên xử lý cao trong kịch bản mô phỏng.

**Ví dụ 2 — vòng đời đơn:** TC-OS-21 tạo đơn, hủy thành công, sau đó PUT status=Đang giao với version mới nhất. Request cuối vẫn thành công. Vì dùng version hợp lệ, phát hiện không phải lỗi stale version: ứng dụng thiếu kiểm tra cạnh chuyển trạng thái.

**Ví dụ 3 — quyền hồ sơ:** TC-USR-2 gửi role=admin qua thao tác cập nhật hồ sơ thường và nhận role đã đổi. TC-USR-1 cũng tái hiện thao tác thiếu xác thực. Cần xem riêng quyền truy cập và lọc trường; có version không thay thế được hai kiểm soát đó.

**Ví dụ 4 — mật khẩu:** TC-USR-4 kiểm tra dữ liệu lưu nội bộ bằng password_audit, không công bố hash. Bản ghi không đáp ứng tiêu chí bcrypt và lưu an toàn. Dấu hiệu phù hợp để xem xét CWE-256 [5]; việc gán CWE trong báo cáo là phân loại đề xuất, không phải CVE được cấp.

Nhóm thiếu phân quyền được đối chiếu CWE-862 [3]; lỗi validation được tham khảo CWE-20 [4]. Mức ưu tiên ở đây là đánh giá định tính theo tác động và khả năng thao tác trong harness, chưa phải điểm CVSS của hệ thống triển khai.
''')
    compare_table = ['| Seed | Dòng baseline | Dòng feedback | Finding baseline | Finding feedback |', '|---|---|---|---|---|']
    for seed in comparison['seeds']:
        pair = {r['mode']: r for r in comparison['runs'] if r['seed'] == seed}
        b, f = pair['mutation_baseline'], pair['coverage_feedback']
        compare_table.append(f"| {seed} | {b['unique_index_lines']} | {f['unique_index_lines']} | {b['finding_occurrences']} | {f['finding_occurrences']} |")
    pages.append('''### 4.3. So sánh hai chiến lược chọn input

Thiết kế dùng cùng 6 target, seed 7/11/42 và 12 input/target/seed. Mỗi chiến lược chạy 216 input, tổng 432. Hai chiến lược đều đo coverage, nhưng baseline luôn đột biến từ seed đầu; feedback dùng coverage mới để chọn cha. Thứ tự chạy được đảo giữa các seed.

''' + '\n'.join(compare_table) + '''

Bảng 5. Độ bao phủ và lượt finding trong thực nghiệm có kiểm soát cấu hình. Không cộng số dòng của các seed thành số dòng duy nhất.

Không có lỗi bộ chạy. Baseline có 120 finding và 20 response HTTP 5xx; feedback có 104 finding và 37 HTTP 5xx. Cả hai bằng coverage ở seed 7 và 42; feedback ít hơn 8 dòng ở seed 11. Chưa quan sát lợi ích coverage feedback trong cấu hình này.

Kết quả khác kỳ vọng đơn giản “dùng coverage luôn tốt hơn”. Một khả năng cần khảo sát tiếp là nhiều dòng mới đến từ setup hoặc nhánh lỗi, trong khi mutation trên input đã biến dạng ít hữu ích cho nghiệp vụ. Đây là giả thuyết giải thích, chưa phải nguyên nhân đã chứng minh bằng thí nghiệm loại trừ.

Ngân sách nhỏ, ba seed và coverage dòng chưa đủ chứng minh ưu thế thống kê. Tác vụ tái hiện chạy đồng thời trong phần đầu nên số giây chỉ mô tả lần chạy, không dùng kết luận nhanh/chậm. Nguồn chi tiết: results/comparison_3seeds/report.json và PhanTich_SoSanh.md.
''')
    pages.append('''## 5. Thảo luận và hạn chế

Fuzzing cung cấp phản ví dụ hữu hạn khi oracle phát hiện vi phạm, nhưng không chứng minh vắng lỗi khi không tìm được input gây lỗi. PASS chỉ nói về mẫu và điều kiện đã kiểm tra. Một oracle yếu có thể bỏ sót lỗi; oracle sai có thể báo nhầm, vì vậy bộ công cụ có 19 kiểm tra hồi quy cho thống kê, giao thức, timeout, version và bộ đếm.

Mô hình datastore có thể khác MongoDB về ép kiểu, ràng buộc và concurrency. Dịch vụ email/notification/closet được mô phỏng; không suy ra độ tin cậy tích hợp thật. Phép đo timing dùng dữ liệu cục bộ và bị ảnh hưởng tải máy. Chưa đo rate limiting, thông lượng, retry khi lỗi hoặc toàn bộ thuộc tính CTL. Các mẫu lỗi bộ nhớ C là bài minh họa riêng.

Các nhóm dấu hiệu cần được điều tra và sửa theo ưu tiên: bảo vệ endpoint và whitelist trường, tính tiền theo catalog, kiểm tra enum/vòng đời, validation trước datastore, hash mật khẩu, rồi xác minh bằng chạy lại cùng input. Các đề xuất này chưa phải bản vá đã được triển khai trong thí nghiệm báo cáo.

Hướng mở rộng là tăng ngân sách và seed, tách coverage setup khỏi request, giảm input gây lỗi, đo trên datastore thử nghiệm gần triển khai và tạo cặp trước/sau bản vá. Không thay đổi mục tiêu thử nghiệm giữa các lần so sánh mà thiếu ghi chú phiên bản.

## 6. Kết luận

Đề tài đã triển khai quy trình đặc tả–testcase–fuzzing–oracle–bằng chứng cho hệ thống bán hàng mô phỏng, thực thi đủ 48 testcase API và tái hiện 21 kịch bản vi phạm đại diện. So sánh 432 input chưa cho thấy lợi thế của lựa chọn theo coverage trong ngân sách đang dùng. Giá trị chính là bằng chứng có thể truy vết và tái hiện, cùng cách phân biệt lỗi ứng dụng, lỗi bộ chạy và giới hạn mô phỏng. Việc xác nhận trên triển khai và đánh giá sau sửa lỗi là công việc tiếp theo.
''')
    references = '''## Tài liệu tham khảo

[1] Vũ Quang Dũng. Hướng dẫn Báo cáo Bài tập lớn: Cấu trúc, định dạng và tiêu chí đánh giá. CSE703093 – An toàn phần mềm, Trường Đại học Phenikaa. Tài liệu do người dùng cung cấp, BTL_02_HuongDan_BaoCao (2).pdf; không có năm xuất bản được xác nhận.

[2] Spec Verification Lab. Hướng dẫn tạo dữ liệu đầu vào bằng file Excel theo template đặc tả. Tài liệu do người dùng cung cấp, Hướng dẫn tạo Spec file dạng excel (2).pdf, 28 trang; không có năm xuất bản được xác nhận.

[3] MITRE. CWE-862: Missing Authorization. https://cwe.mitre.org/data/definitions/862.html. Truy cập 15/10/2026.

[4] MITRE. CWE-20: Improper Input Validation. https://cwe.mitre.org/data/definitions/20.html. Truy cập 15/10/2026. Dùng tham khảo khái niệm chung, chưa gán mọi lỗi vào một mã gốc duy nhất.

[5] MITRE. CWE-256: Plaintext Storage of a Password. https://cwe.mitre.org/data/definitions/256.html. Truy cập 15/10/2026.

[6] Node.js. Inspector API. https://nodejs.org/api/inspector.html. Truy cập 15/10/2026; phiên bản thực nghiệm cần đối chiếu package và môi trường chạy, không mặc định trùng phiên bản trang tài liệu hiện hành.

[7] Mã nguồn và kết quả thực nghiệm trong dự án nộp kèm; phiên bản code/workbook được xác định bằng source_sha256 trong report.json và SHA-256 trong spec_snapshot.json.

## Phụ lục: cấu hình, tái hiện và bằng chứng

```powershell
python -m pip install -r fuzzing-security-testing/requirements.txt
npm ci --prefix fuzzing-security-testing/node_harness
python -m unittest discover -s fuzzing-security-testing/tests -p test_campaign.py
python fuzzing-security-testing/fuzzers/run_ch4.py --seed 7 --fuzz-cases 8 --output fuzzing-security-testing/results/reproduction_48
python fuzzing-security-testing/fuzzers/controlled_comparison.py --iterations-per-target 12 --seeds 7 11 42 --output fuzzing-security-testing/results/comparison_new
```

| Bằng chứng | Đường dẫn từ thư mục fuzzing-security-testing |
|---|---|
| Đặc tả và testcase | BaoCaoATPM/Module_*_spec.xlsx |
| Kết quả 48 testcase | results/spec_completed_48/report.json |
| Bảng kết quả Excel | results/spec_completed_48/KetQuaKiemThu.xlsx |
| 42 lượt tái hiện | results/reproduced_48/evidence/ |
| Nhóm dấu hiệu | results/reproduced_48/PhanNhom_PhatHien.md |
| So sánh 3 seed | results/comparison_3seeds/report.json |
| Mã nguồn chính | fuzzers/ và node_harness/; helper server/concurrency.js ở thư mục gốc dự án |

Khi nộp cần giữ đúng cấu trúc đường dẫn để harness tìm được helper dùng chung. Tệp Word là bản biên soạn; trước khi xuất PDF cần cập nhật mục lục/số trang và đối chiếu hình, bảng, đường dẫn với bộ source/results nộp kèm.

### Ghi nhận hỗ trợ AI

AI hỗ trợ rà soát/viết lại đặc tả, tổ chức và chỉnh sửa code fuzzing, bổ sung testcase/oracle, điều phối kiểm tra, tổng hợp bằng chứng và biên soạn báo cáo. Người thực hiện cần tự kiểm tra, hiểu mã nguồn và chịu trách nhiệm về kết luận; không trình bày các phần hỗ trợ này là hoàn toàn tự viết không có hỗ trợ.
'''
    cover = '''# PHƯƠNG PHÁP KIỂM THỬ THEO FUZZING CHO HỆ THỐNG BÁN HÀNG

TRƯỜNG ĐẠI HỌC PHENIKAA

KHOA HỆ THỐNG THÔNG TIN

BÁO CÁO BÀI TẬP LỚN — CSE703093: AN TOÀN PHẦN MỀM

Giảng viên: ThS. Vũ Quang Dũng

Họ tên: [ĐIỀN HỌ TÊN]

Mã số sinh viên: [ĐIỀN MSSV]

Lớp: [ĐIỀN LỚP]

Học kỳ: [ĐIỀN HỌC KỲ]

Bản biên soạn có bằng chứng thực nghiệm; cần điền thông tin, kiểm tra số trang và xuất PDF trước khi nộp.
'''
    formal_path = ROOT / 'results/formal_verification/report.json'
    if formal_path.exists():
        formal = json.loads(formal_path.read_text(encoding='utf-8'))
        if formal['errors']:
            raise ValueError('Formal verification contains runner errors')
        pages = [page.replace('19 kiểm tra hồi quy', '26 kiểm tra hồi quy')
                 .replace('Chưa đo rate limiting, thông lượng, retry khi lỗi hoặc toàn bộ thuộc tính CTL.',
                          'Chưa đo rate limiting, thông lượng và retry khi lỗi. CTL/SMT đã được kiểm tra riêng trên mô hình workbook, không phải chứng minh mọi hành vi API.')
                 for page in pages]
        references += '''

### Phụ lục B. Xác minh mô hình đặc tả

Nguồn: results/formal_verification/report.json và report.md. Bộ xác minh do dự án
cài đặt theo quy tắc PDF, không phải core.spec_pipeline gốc vì gói đó chưa có trong
các toolkit được cung cấp. Không khẳng định đã chạy parser/pipeline gốc của giảng viên.

Năm workbook được kiểm tra cấu trúc, ID, nhãn và cạnh chuyển; không có deadlock.
Có 24 công thức CTL: 12 đúng, 12 sai theo ngữ nghĩa thỏa tại tất cả trạng thái khởi tạo.
Các phản ví dụ và trạng thái khởi tạo không thỏa được lưu đầy đủ trong JSON.

Năm công thức AG(!violation) và năm cặp mutex chính sai do conflict được cấy theo
template. Công thức giữ trạng thái actual_completed sai vì mô hình cho phép quay lui.
Sáu công thức EF sai ở một số trạng thái khởi tạo: các cụm nghiệp vụ/khóa là thành phần
rời nhau. Điều này khác với khẳng định không có bất kỳ đường nào tới main_locked;
cần đọc failing_initial_states và chọn mô hình thành phần nếu muốn kiểm chứng riêng.
Không đổi nhãn hay tập khởi tạo chỉ để biến kết quả thành PASS.

Có 25 kiểm tra SMT: SAT điều kiện nghiệp vụ và SAT overflow int32 được chạy bằng
solver riêng; còn có phép kiểm tra kết hợp để tránh suy luận nhầm từ hai kết quả độc lập.
Mô hình số học không chứng minh JavaScript dùng int32. Z3 và GCC được ghi phiên bản
trong report.json để tái hiện.

Mã sinh minh họa có 20 trường hợp lỗi C được cảnh báo: 5 pattern strcpy tiềm năng,
15 trường hợp leak/double-free/use-after-free bởi GCC -fanalyzer. Không có xác nhận ASan.
Mười hàm Python được thực thi trên 70 input: divide ở biên và magic_value bằng input
Z3 cho nhánh mẫu cố định; không nhận là bộ concolic tổng quát. Các lỗi cấy này không
được cộng vào 21 testcase API FAIL.

```powershell
python -m pip install -r fuzzing-security-testing/supporting-analysis/requirements.txt
python -m unittest discover -s fuzzing-security-testing/tests -p "test_*.py"
python fuzzing-security-testing/supporting-analysis/tools/verify_spec_workbooks.py
```

Máy cần GCC có -fanalyzer. Nếu thiếu compiler hoặc solver, công cụ ghi ERROR thay
vì giả định kiểm chứng thành công. Kết quả mới của API vẫn là 27 PASS/21 FAIL trên
48 testcase; thêm kiểm chứng mô hình không thay đổi các số liệu API đó.
'''
    return cover, pages, references


def export_word(cover, pages, references, target):
    from docx import Document
    from docx.shared import Pt, Cm
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(2)
    section.left_margin, section.right_margin = Cm(2.5), Cm(2)
    for name in ['Normal', 'Heading 1', 'Heading 2', 'Heading 3']:
        style = document.styles[name]
        style.font.name = 'Times New Roman'
        style.font.size = Pt(12)
        style.paragraph_format.line_spacing = 1.5
        style.paragraph_format.space_after = Pt(6)
        style.element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')

    def render(text):
        lines = text.strip().splitlines()
        code = False
        table = None
        for line in lines:
            if line.startswith('```'):
                code = not code
                table = None
                continue
            if code:
                p = document.add_paragraph()
                run = p.add_run(line)
                run.font.name = 'Consolas'
                run.font.size = Pt(9)
                p.paragraph_format.line_spacing = 1
                p.paragraph_format.space_after = Pt(0)
                continue
            if line.startswith('|'):
                cells = [part.strip() for part in line.strip('|').split('|')]
                if all(set(value) <= set('-: ') for value in cells):
                    continue
                if table is None:
                    table = document.add_table(rows=0, cols=len(cells))
                    table.style = 'Table Grid'
                row = table.add_row()
                for cell, value in zip(row.cells, cells):
                    cell.text = value
                    for p in cell.paragraphs:
                        p.paragraph_format.line_spacing = 1
                        p.paragraph_format.space_after = Pt(4)
                        for run in p.runs:
                            run.font.size = Pt(10)
                continue
            table = None
            if not line:
                continue
            if line.startswith('## Phụ lục:'):
                document.add_page_break()
            if line.startswith('#'):
                level = min(3, len(line) - len(line.lstrip('#')))
                document.add_heading(line.lstrip('#').strip(), level=level)
            else:
                content = line[2:] if line.startswith('- ') else line
                document.add_paragraph(content.replace('**', ''), style='List Bullet' if line.startswith('- ') else None)

    for line in cover.strip().splitlines():
        if line.strip():
            p = document.add_paragraph(line.lstrip('#').strip())
            p.alignment = 1
            if line.startswith('#'):
                for run in p.runs:
                    run.bold = True
                    run.font.size = Pt(16)
    document.add_page_break()
    document.add_paragraph('MỤC LỤC', style='Title')
    p = document.add_paragraph()
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), 'TOC \\o "1-3" \\h \\z \\u')
    p._p.append(field)
    document.add_paragraph('Trong Word: nhấn Ctrl+A rồi F9 để cập nhật mục lục và số trang trước khi xuất PDF.')
    for page in pages:
        document.add_page_break()
        render(page)
    document.add_page_break()
    render(references)
    footer = section.footer.paragraphs[0]
    footer.alignment = 1
    page_number = OxmlElement('w:fldSimple')
    page_number.set(qn('w:instr'), 'PAGE')
    footer._p.append(page_number)
    setting = OxmlElement('w:updateFields')
    setting.set(qn('w:val'), 'true')
    document.settings.element.append(setting)
    document.save(target)
    check = Document(target)
    assert len(check.tables) == 6, len(check.tables)
    assert len(check.sections) == 1


def main():
    cover, pages, references = compile_report()
    path = ROOT / 'BaoCaoATPM/BaoCao_Fuzzing_DuThao.md'
    backup = ROOT / 'BaoCaoATPM/SaoLuu_TruocLamRo' / ('bao_cao_truoc_dinh_dang_' + datetime.now().strftime('%Y%m%d_%H%M%S'))
    backup.mkdir(parents=True)
    for existing in (path, path.with_suffix('.docx'), path.with_suffix('.pdf')):
        if existing.exists():
            shutil.copy2(existing, backup / existing.name)
    path.write_text('\n\n<!-- PAGE_BREAK -->\n\n'.join([cover, *pages, references]), encoding='utf-8')
    target = path.with_suffix('.docx')
    export_word(cover, pages, references, target)
    print(target)
    print('Word đã tạo; cần cập nhật mục lục và kiểm tra số trang sau khi render.')


if __name__ == '__main__':
    main()
