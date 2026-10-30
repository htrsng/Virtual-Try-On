"""Build auditable Vietnamese test minutes from saved evidence; no campaign rerun."""
import sys, json, hashlib
from pathlib import Path
from collections import Counter
from datetime import datetime, timezone, timedelta

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT.parent / '.codex-report-deps'))
import openpyxl
from docx import Document
from docx.shared import Pt, Cm

OUT = ROOT / 'BaoCaoATPM' / 'BienBan_Test_20261030'
OUT.mkdir(exist_ok=True)
def read(rel):
    return json.loads((ROOT / rel).read_text(encoding='utf-8'))
def sha(p):
    return hashlib.sha256(p.read_bytes()).hexdigest()
report = read('results/spec_excel_aligned/report.json')
snapshot = read('results/spec_excel_aligned/spec_snapshot.json')
black = read('results/blackbox/e97d715116e3/report.json')
traces = {t['testcase']['ID']: t for t in report['traceability']}
old = {t['ID']: t for s in snapshot for t in s['sheets']['TestCase']}
specs, current, manifest = [], [], []
for p in sorted((ROOT / 'BaoCaoATPM').glob('Module_*_spec.xlsx')):
    w = openpyxl.load_workbook(p, data_only=True)
    sheets = {s.title: list(s.values) for s in w}
    rows = [dict(zip(sheets['TestCase'][0], r)) for r in sheets['TestCase'][1:] if r[0]]
    assert len({r['ID'] for r in rows}) == len(rows)
    current.extend(rows)
    previous = next(s for s in snapshot if s['file'] == p.name)
    specs.append((p.name, rows, sha(p), previous['sha256']))
    manifest.append({'path': str(p.relative_to(ROOT)), 'sha256': sha(p), 'snapshot_sha256': previous['sha256'], 'sheet_rows': {s: len([r for r in rs[1:] if r and r[0] is not None]) for s, rs in sheets.items()}})
    w.close()
assert len(current) == 134 and len(old) == 104
assert Counter(e['status'] for e in report['executions']) == report['execution_counts']
assert sum(black['status_counter'].values()) == black['iterations']
assert len(black['errors']) == 240 and len(black['vuln_findings']) == 1669

lines = []
def p(t=''): lines.append(t)
def table(headers, rows):
    p('| ' + ' | '.join(headers) + ' |'); p('| ' + ' | '.join(['---']*len(headers)) + ' |')
    for row in rows: p('| ' + ' | '.join(str(c).replace('|','/').replace('\n','; ') for c in row) + ' |')
    p()
p('# BIÊN BẢN KIỂM THỬ AN TOÀN PHẦN MỀM')
p('Số: BBKT-20261030-01 · Ngày lập: 30/10/2026 · Hệ thống bán hàng mô phỏng trong Virtual-Try-On-sync.\n')
p('## 1. Thông tin và căn cứ')
p('Đơn vị/nhóm thực hiện: ................................; Người kiểm thử: ................................; Người xác nhận: ................................\n')
p('Biên bản được lập bằng cách đọc 5 workbook Excel (đủ 10 sheet/workbook), mã nguồn bộ chạy/oracle và bằng chứng JSON đã lưu. Đây là biên bản tổng hợp kết quả lịch sử, không phải xác nhận đã chạy mới ngày 30/10/2026. Không điền giả tên người thực hiện, thời gian chạy hoặc chữ ký.\n')
stamp = datetime.fromisoformat(report['created_utc']).astimezone(timezone(timedelta(hours=7)))
p(f'Campaign đối chiếu Excel: results/spec_excel_aligned/report.json; thời điểm tạo báo cáo {stamp:%d/%m/%Y %H:%M:%S} (UTC+7); seed {report["random_seed"]}; Python được ghi trong báo cáo: {report["python_version"]}. Môi trường: node_harness_stub; Express xử lý request thực, datastore và dịch vụ là mô phỏng. Không đại diện cho production/MongoDB thật.\n')
p('## 2. Đối chiếu đặc tả và phạm vi')
table(['Workbook','TC hiện tại','Có kết quả lịch sử','Chưa có kết quả','PASS/FAIL lịch sử'],[(name,len(rows),sum(r['ID'] in traces for r in rows),sum(r['ID'] not in traces for r in rows),str(dict(Counter(traces[r['ID']]['status'] for r in rows if r['ID'] in traces)))) for name,rows,_,_ in specs])
p('Excel hiện tại có 134 testcase; snapshot lúc chạy có 104. Có 30 ID bổ sung, 75 ID trong snapshot chưa thực thi, tổng cộng 105 testcase hiện tại chưa có kết quả campaign API này. README ghi 104 là số của snapshot, không phải số đếm hiện tại.\n')
p('Cả 5 SHA-256 workbook hiện tại đều khác snapshot. Đối chiếu từng trường trong sheet TestCase của 104 ID chung không phát hiện thay đổi nội dung; 30 ID bổ sung chưa có kết quả. Khác mã băm không tự chứng minh 104 testcase cũ đã đổi yêu cầu. Kết quả PASS/FAIL dưới đây chỉ thuộc kịch bản/oracle và snapshot lúc chạy; không xác nhận toàn bộ workbook hiện tại. Phụ lục Thay_doi_dac_ta chỉ có hàng tiêu đề vì không có trường TestCase nào thay đổi trên các ID chung.\n')
p('## 3. Kết quả campaign API theo đặc tả')
p('29 testcase đã thực thi, 9 PASS, 20 FAIL; 59 lượt kịch bản gồm 19 PASS và 40 FAIL; không có ERROR trong báo cáo này. Tỷ lệ PASS theo testcase đã chạy: 31,03%; mức hiện diện bằng chứng theo ID trên Excel hiện tại: 29/134 = 21,64%, không phải tỷ lệ bao phủ yêu cầu hoặc mã nguồn. Một lượt kịch bản có thể chứa nhiều HTTP request.\n')
table(['ID','Tên tại thời điểm chạy','Kết quả lịch sử','Lượt'],[(t['testcase']['ID'],t['testcase']['TieuDe'],t['status'],t['executions']) for t in report['traceability']])
p('PASS chỉ xác nhận các mẫu và điều kiện oracle đã cài đặt. Ví dụ TC-AUTH-2 có 5 lượt trong campaign, không xác nhận đã lặp 2.000 lần như bước Excel mô tả; không tự bao phủ mọi điều kiện bảo mật mở rộng.\n')
p('## 4. Bằng chứng cho các testcase FAIL')
for t in report['traceability']:
    if t['status'] != 'FAIL': continue
    tc = t['testcase']; executions = [e for e in report['executions'] if e['case']['testcase']==tc['ID'] and e['status']=='FAIL']
    p(f'### {tc["ID"]} — {tc["TieuDe"]}')
    p('Mong đợi theo snapshot: '+tc['KetQuaMongDoi']+'\n')
    table(['Case ID','Oracle thực tế','HTTP quan sát','Tham chiếu'],[(e['case']['id'],json.dumps(e['case']['check'],ensure_ascii=False),', '.join(str(o['status']) for o in e['evidence']['observations']),f"report.json → executions[{report['executions'].index(e)}]") for e in executions])
    e=executions[0]; ev=e['evidence']; obs=ev['observations']
    # Include diagnostic fields, never tokens or full response bodies.
    messages=[{k:o['body'][k] for k in ('message','error','details','status','role') if isinstance(o.get('body'),dict) and k in o['body']} for o in obs]
    p('Trích quan sát: '+json.dumps(messages,ensure_ascii=False)+'\n')
    if ev.get('password_audit'): p('Kiểm tra lưu mật khẩu: '+json.dumps(ev['password_audit'],ensure_ascii=False)+'\n')
p('Các mục FAIL là nhóm theo testcase; chưa có phân tích chứng minh đây là 20 lỗi gốc độc lập. Dữ liệu request, actor, entity_before/after (nếu được lưu), response và toàn bộ chuỗi thao tác tra trong executions của JSON gốc theo Case ID ở trên.\n')
p('## 5. Blackbox 3.500 input (campaign riêng)')
p(f'Nguồn: results/blackbox/e97d715116e3/report.json. Seed 42, 3.500 input, thời gian được ghi {black["elapsed_seconds"]:.3f} giây. 240 HTTP 5xx (6,86%), 1.669 lượt finding nghiệp vụ/bảo mật. Không dùng 1.669 làm số lỗi gốc; không cộng với 20 testcase FAIL.\n')
table(['Target','HTTP','Số lượt'],[(k.rsplit(':',1)[0],k.rsplit(':',1)[1],v) for k,v in sorted(black['status_counter'].items())])
table(['Target có 5xx','Số lượt'],Counter(e['target'] for e in black['errors']).items())
p('Ví dụ iteration 1: tạo đơn thiếu totalAmount trả HTTP 500 với thông báo required; finding iteration 3 ghi sửa hồ sơ bản thân không có token vẫn HTTP 200; iteration 4 ghi sửa hồ sơ người khác bằng token user thường vẫn HTTP 200. Đây là quan sát trên harness; 500 không chứng minh tiến trình crash.\n')
p('## 6. Demo coverage, concolic và khảo sát biên')
p('results/demo_chuong4_final/: blackbox 42 input, 1 lỗi HTTP ứng dụng, 22 quan sát nghiệp vụ/bảo mật; khảo sát coverage 16 input, 6 lỗi HTTP ứng dụng và 6 quan sát. Coverage-guided là greybox; không gọi toàn bộ giai đoạn này là symbolic execution.\n')
p('Concolic riêng cho totalAmount: 5 input, 4 input replay, 4 input vi phạm, 2 input HTTP 5xx, 1 truy vấn UNSAT. number trả 200; null và object trả 500; numeric_string và boolean trả 200 nhưng bị phân loại chấp nhận sai kiểu. Mô hình thủ công total_kind trong [0,7], Z3 phiên bản ghi nhận 5.1.0; không chứng minh toàn bộ JavaScript hoặc toàn hệ thống. requirements.txt hiện quy định z3-solver<5 khác phiên bản ghi trong bằng chứng: cần tái lập môi trường trước khi so sánh lần chạy mới.\n')
p('results/boundary_smoke_final/: blackbox 14 input (1 lỗi HTTP, 8 quan sát); khảo sát coverage 4 input (4 lỗi HTTP). Probe thay productId và độ dài name trong [0,16] trong khi shippingFee sai kiểu được giữ nguyên; các mẫu đều 5xx, không có transition, complete=false. Vì vậy chưa xác định được biên lỗi của productId/name, chưa loại trừ tác động che khuất từ shippingFee. Không suy ra toàn miền [0,16] đã được kiểm tra đầy đủ.\n')
p('## 7. Sai lệch và giới hạn bằng chứng')
p('Hai lần thử ngày 28/10/2026 được Kiem_tra_runtime.json ghi SETUP_ERROR: harness chưa sẵn sàng sau 15 giây; nguyên nhân chưa xác định. Đây không phải campaign hoàn tất, không cộng thành PASS/FAIL của ứng dụng.\n')
p('Chưa có bằng chứng xác nhận crash tiến trình, hang hoặc memory corruption bằng sanitizer; giá trị null trong mục confirmed_* là chưa đo, không phải số 0. Chưa xác minh đầy đủ CTL/SMT, concurrency, timing, rate limit, tích hợp email/virtual closet và 105 testcase còn lại. Năm sheet mô hình/thuộc tính hay mã sinh trong Excel là đặc tả, không tự trở thành kết quả chạy.\n')
p('## 8. Kết luận và xử lý đề nghị')
p('Kết luận theo bằng chứng lịch sử: campaign API chưa đạt, vì 20/29 testcase vi phạm oracle; chưa đủ cơ sở xác nhận an toàn hoặc nghiệm thu toàn bộ 134 testcase hiện tại. Chưa có kết quả retest chứng minh các vi phạm đã được sửa.\n')
p('Ưu tiên sửa kiểm tra xác thực/quyền trên tài khoản, sản phẩm và trạng thái đơn; ngăn mass assignment role, lưu mật khẩu bằng hash phù hợp; kiểm tra kiểu/miền đầu vào trước tầng lưu trữ; tính giá/tổng từ dữ liệu server; bảo vệ trạng thái kết thúc và định danh sản phẩm; đồng nhất phản hồi đăng nhập. Sau sửa cần chạy lại theo Excel hiện tại, bổ sung oracle cho điều kiện mới, lưu snapshot/hash/runtime và retest từng FAIL cùng testcase hồi quy. Mức ưu tiên là đề nghị xử lý, chưa phải chấm điểm CVSS.\n')
p('## 9. Nguồn và truy vết')
sources=['results/spec_excel_aligned/report.json','results/spec_excel_aligned/spec_snapshot.json','results/spec_excel_aligned/excluded_testcases.json','results/blackbox/e97d715116e3/report.json','results/demo_chuong4_final/report.json','results/demo_chuong4_final/concolic/report.json','results/boundary_smoke_final/report.json','results/boundary_smoke_final/whitebox/bounded_analysis.json','BaoCaoATPM/BaoCao_Trang/Kiem_tra_runtime.json','fuzzers/oracles.py','fuzzers/spec_campaign.py','README.md','requirements.txt']
for rel in sources:
    f=ROOT/rel; manifest.append({'path':rel,'sha256':sha(f)})
table(['Nguồn (tương đối từ fuzzing-security-testing)','SHA-256 hiện tại'],[(m['path'],m['sha256']) for m in manifest])
p('Phụ lục Excel đi kèm: toàn bộ 134 testcase hiện tại; kết quả lịch sử được tách riêng; thay đổi từng trường so với snapshot; 59 lượt kịch bản cùng oracle và HTTP; mã băm nguồn. Không có thao tác chạy campaign trong lần lập biên bản này.\n')
p('## 10. Xác nhận')
table(['Người lập biên bản','Đại diện kiểm thử','Đại diện phụ trách hệ thống'],[['(Ký, ghi rõ họ tên)','(Ký, ghi rõ họ tên)','(Ký, ghi rõ họ tên)'],['........................','........................','........................']])
(OUT/'Bien_ban_test.md').write_text('\n'.join(lines),encoding='utf-8')
(OUT/'Manifest_bang_chung.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2),encoding='utf-8')

wb=openpyxl.Workbook(); ws=wb.active; ws.title='Doi_chieu_134_TC'
ws.append(['Workbook','ID','Tiêu đề hiện tại','Mong đợi hiện tại','Kết quả lịch sử (snapshot)','Lượt lịch sử','Tình trạng phiên bản','Trường thay đổi'])
diffs=[]
for name,rows,_,_ in specs:
    for tc in rows:
        tid=tc['ID']; changed=[k for k in tc if tid in old and tc[k]!=old[tid].get(k)]
        ws.append([name,tid,tc['TieuDe'],tc['KetQuaMongDoi'],traces.get(tid,{}).get('status','CHƯA CÓ KẾT QUẢ'),traces.get(tid,{}).get('executions',0),'ID bổ sung sau snapshot' if tid not in old else 'So theo ID; hash workbook khác',', '.join(changed)])
        for k in changed: diffs.append([tid,k,old[tid].get(k),tc[k]])
ws=wb.create_sheet('Thay_doi_dac_ta');ws.append(['ID','Trường','Snapshot lúc chạy','Excel hiện tại'])
for row in diffs:ws.append(row)
ws=wb.create_sheet('59_luot_lich_su');ws.append(['Case ID','Testcase','Kết quả','Oracle','HTTP theo thứ tự','Tham chiếu JSON'])
for i,e in enumerate(report['executions']):ws.append([e['case']['id'],e['case']['testcase'],e['status'],json.dumps(e['case']['check'],ensure_ascii=False),', '.join(str(o['status']) for o in e.get('evidence',{}).get('observations',[])),f'executions[{i}]'])
ws=wb.create_sheet('Nguon_SHA256');ws.append(['Nguồn','SHA256 hiện tại','SHA256 snapshot'])
for m in manifest:ws.append([m['path'],m['sha256'],m.get('snapshot_sha256','')])
for ws in wb:
    ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
    for cell in ws[1]:cell.font=openpyxl.styles.Font(bold=True,color='FFFFFF');cell.fill=openpyxl.styles.PatternFill('solid',fgColor='24476A')
    for col in ws.columns:ws.column_dimensions[col[0].column_letter].width=min(65,max(18,max(len(str(c.value or '')) for c in col)+2))
    for row in ws.iter_rows(min_row=2):
        for c in row:c.alignment=openpyxl.styles.Alignment(vertical='top',wrap_text=True)
wb.save(OUT/'Phu_luc_doi_chieu_134_testcase.xlsx')

doc=Document();sec=doc.sections[0];sec.top_margin=sec.bottom_margin=Cm(1.8)
style=doc.styles['Normal'];style.font.name='Times New Roman';style.font.size=Pt(11)
i=0
while i<len(lines):
    line=lines[i]
    if line.startswith('| '):
        group=[]
        while i<len(lines) and lines[i].startswith('| '):group.append(lines[i]);i+=1
        rows=[[c.strip() for c in x.strip('|').split('|')] for j,x in enumerate(group) if j!=1]
        t=doc.add_table(rows=0,cols=len(rows[0]));t.style='Table Grid'
        for row in rows:
            for c,value in zip(t.add_row().cells,row):c.text=value
        continue
    if line.startswith('# '):doc.add_heading(line[2:],0)
    elif line.startswith('## '):doc.add_heading(line[3:],1)
    elif line.startswith('### '):doc.add_heading(line[4:],2)
    elif line:doc.add_paragraph(line)
    i+=1
doc.save(OUT/'Bien_ban_test.docx')
assert openpyxl.load_workbook(OUT/'Phu_luc_doi_chieu_134_testcase.xlsx',read_only=True)['Doi_chieu_134_TC'].max_row==135
assert len(Document(OUT/'Bien_ban_test.docx').tables)>20
print(json.dumps({'output':str(OUT),'testcases_current':len(current),'testcases_snapshot':len(old),'changed_fields':len(diffs),'files':[f.name for f in OUT.iterdir()]},ensure_ascii=False))
